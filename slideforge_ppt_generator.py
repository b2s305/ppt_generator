
from __future__ import annotations

import colorsys
import datetime
import math
import os
import random
import re
import subprocess
import sys
import threading
import time

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
    from lxml import etree
except ImportError:  # friendly message instead of a traceback
    print("Missing dependency.  Run:  pip install python-pptx")
    try:
        import tkinter as _tk
        from tkinter import messagebox as _mb
        _r = _tk.Tk()
        _r.withdraw()
        _mb.showerror("Missing package", "Please run:\n\npip install python-pptx")
    except Exception:
        pass
    sys.exit(1)


# =============================================================================
#  COLOUR UTILITIES
# =============================================================================
def hsl(h, s, l):
    r, g, b = colorsys.hls_to_rgb((h % 360) / 360.0, max(0, min(1, l)), max(0, min(1, s)))
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def rgb(hexs):
    return RGBColor.from_string(hexs)


def _lin(c):
    return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4


def luminance(hexs):
    r, g, b = [int(hexs[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    return 0.2126 * _lin(r) + 0.7152 * _lin(g) + 0.0722 * _lin(b)


def contrast(a, b):
    la, lb = luminance(a), luminance(b)
    if la < lb:
        la, lb = lb, la
    return (la + 0.05) / (lb + 0.05)


def on_color(bg, light="FFFFFF", dark="0F172A"):
    """Readable text colour for a given background."""
    return light if contrast(bg, light) >= contrast(bg, dark) else dark


def mix(a, b, t):
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "%02X%02X%02X" % tuple(round(x + (y - x) * t) for x, y in zip(ca, cb))


# =============================================================================
#  LOOK GENERATOR  (palette + style + fonts)  -> different every time
# =============================================================================
HARMONY = {
    "analogous": (30, -30),
    "complementary": (180, 150),
    "triadic": (120, 240),
    "split-complementary": (150, 210),
    "tetradic": (90, 180),
}

PRESETS = {
    "🎲  Surprise me (new every time)": None,
    "Ocean Breeze": (205, "analogous", "light"),
    "Midnight Neon": (265, "triadic", "dark"),
    "Sunset Glow": (14, "split-complementary", "light"),
    "Forest Calm": (150, "analogous", "light"),
    "Royal Gold": (42, "complementary", "dark"),
    "Berry Pop": (322, "complementary", "light"),
    "Cyber Teal": (178, "split-complementary", "dark"),
    "Crimson Executive": (352, "analogous", "dark"),
}

FONT_PAIRS = {
    "Auto (random pair)": None,
    "Modern  - Segoe UI": ("Segoe UI Semibold", "Segoe UI"),
    "Classic - Georgia + Calibri": ("Georgia", "Calibri"),
    "Clean   - Calibri": ("Calibri", "Calibri"),
    "Bold    - Arial Black + Arial": ("Arial Black", "Arial"),
    "Friendly - Trebuchet MS": ("Trebuchet MS", "Trebuchet MS"),
}
# average glyph width as a fraction of the font size (used for text fitting)
FONT_CW = {"Segoe UI Semibold": 0.55, "Segoe UI": 0.53, "Georgia": 0.56, "Calibri": 0.50,
           "Arial Black": 0.68, "Arial": 0.55, "Trebuchet MS": 0.54}


def make_palette(rng, preset=None, mode="auto"):
    if preset:
        h, harmony, pmode = preset
        h += rng.uniform(-10, 10)
        if mode == "auto":
            mode = pmode
    else:
        h = rng.uniform(0, 360)
        harmony = rng.choice(list(HARMONY))
    dark = (rng.random() < 0.45) if mode == "auto" else (mode == "dark")
    o1, o2 = HARMONY[harmony]
    h2, h3 = h + o1, h + o2
    if dark:
        bg = hsl(h, rng.uniform(.30, .50), rng.uniform(.075, .115))
        surface = hsl(h, rng.uniform(.25, .42), rng.uniform(.15, .19))
        text, muted = "F8FAFC", hsl(h, .18, .74)
        primary = hsl(h, rng.uniform(.70, .92), rng.uniform(.58, .66))
        secondary = hsl(h2, rng.uniform(.65, .90), rng.uniform(.58, .68))
        accent = hsl(h3, rng.uniform(.75, .95), rng.uniform(.60, .70))
    else:
        bg = hsl(h, rng.uniform(.25, .55), rng.uniform(.955, .975))
        surface = "FFFFFF"
        text, muted = hsl(h, rng.uniform(.30, .50), rng.uniform(.11, .16)), hsl(h, .15, .38)
        primary = hsl(h, rng.uniform(.60, .85), rng.uniform(.36, .46))
        secondary = hsl(h2, rng.uniform(.55, .80), rng.uniform(.42, .52))
        accent = hsl(h3, rng.uniform(.70, .95), rng.uniform(.48, .58))
    hero1 = hsl(h, rng.uniform(.55, .80), rng.uniform(.12, .19))
    hero2 = hsl(h2, rng.uniform(.50, .75), rng.uniform(.25, .33))
    return dict(dark=dark, harmony=harmony, bg=bg, surface=surface, text=text, muted=muted,
                primary=primary, secondary=secondary, accent=accent, hero1=hero1, hero2=hero2)


def make_look(seed, theme=None, mode="auto", font=None):
    rng = random.Random(seed)
    pal = make_palette(rng, PRESETS.get(theme), mode)
    style = dict(
        header=rng.choice(["bar", "underline", "band", "tab"]),
        decor=rng.choice(["circles", "rings", "hex", "triangles", "diamonds", "dots"]),
        marker=rng.choice(["num", "dot", "diamond", "gem"]),
        title=rng.choice(["left", "center", "split"]),
        radius=rng.choice([0.08, 0.18, 0.30]),
        angle=rng.choice([20, 35, 60, 120, 135, 160]),
    )
    fonts = FONT_PAIRS.get(font) or rng.choice([v for v in FONT_PAIRS.values() if v])
    return dict(seed=seed, pal=pal, style=style, fonts=fonts)


def describe_look(look):
    p, s = look["pal"], look["style"]
    return ("%s harmony  |  %s mode  |  %s header  |  %s shapes  |  %s"
            % (p["harmony"].title(), "Dark" if p["dark"] else "Light", s["header"],
               s["decor"], look["fonts"][0]))


# =============================================================================
#  TEXT PARSER
# =============================================================================
BULLET_RE = re.compile(r"^(\s*)([-*•–▪●◦>]|\d{1,2}[.)])\s+(.*)$")
HEAD_RE = re.compile(r"^\s*#{1,6}\s*(.+?)\s*#*\s*$")
SLIDE_RE = re.compile(r"^\s*(?:slide|section)\s*\d*\s*[:\-–—]\s*(.+)$", re.I)
DIRECT_RE = re.compile(r"^\s*(title|subtitle|author|presenter|date)\s*[:=]\s*(.+)$", re.I)
NUM_RE = re.compile(r"[\$₹€£]?\d[\d,]*(?:\.\d+)?\s?(?:%|x\b|X\b|[kKMBL]\b|\+|Cr\b)")


def _clean(t):
    return t.replace("**", "").replace("`", "").strip()


def parse_points(raw):
    deck = dict(title=None, subtitle=None, author=None, date=None, slides=[])
    body = []
    for line in raw.replace("\r", "").replace("\t", "    ").split("\n"):
        line = line.rstrip()
        m = DIRECT_RE.match(line)
        if m and not BULLET_RE.match(line):
            k, v = m.group(1).lower(), _clean(m.group(2))
            key = "author" if k == "presenter" else k
            if not deck.get(key):
                deck[key] = v
                continue
        body.append(line)

    def new_slide(title):
        sl = dict(title=_clean(title), raw=[])
        deck["slides"].append(sl)
        return sl

    def add_line(sl, line):
        b = BULLET_RE.match(line)
        if b:
            sl["raw"].append((len(b.group(1)), _clean(b.group(3))))
        else:
            sl["raw"].append((len(line) - len(line.lstrip()), _clean(line)))

    heading_mode = any(HEAD_RE.match(l) or SLIDE_RE.match(l) for l in body)
    if heading_mode:
        cur = None
        for line in body:
            if not line.strip():
                continue
            h = HEAD_RE.match(line) or SLIDE_RE.match(line)
            if h:
                cur = new_slide(h.group(1))
                continue
            if cur is None:
                if not deck["title"] and not BULLET_RE.match(line):
                    deck["title"] = _clean(line)
                    continue
                cur = new_slide("Overview")
            add_line(cur, line)
    else:
        blocks, cur = [], []
        for line in body:
            if line.strip():
                cur.append(line)
            elif cur:
                blocks.append(cur)
                cur = []
        if cur:
            blocks.append(cur)
        if len(blocks) == 1:
            lines = blocks[0]
            if not deck["title"] and not BULLET_RE.match(lines[0]) and len(lines) > 1:
                deck["title"] = _clean(lines[0])
                lines = lines[1:]
            sl = new_slide("Key Points")
            for l in lines:
                add_line(sl, l)
        else:
            for i, blk in enumerate(blocks):
                first = blk[0]
                if not BULLET_RE.match(first):
                    if i == 0 and len(blk) == 1 and not deck["title"]:
                        deck["title"] = _clean(first)
                        continue
                    sl = new_slide(first.strip().rstrip(":"))
                    rest = blk[1:]
                else:
                    sl = new_slide("Part %d" % (i + 1))
                    rest = blk
                for l in rest:
                    add_line(sl, l)

    # normalise bullet levels (0 = main point, 1 = sub-point)
    for sl in deck["slides"]:
        raw = sl.pop("raw")
        base = min([i for i, _ in raw], default=0)
        sl["points"] = [(1 if (i - base) >= 2 else 0, t) for i, t in raw if t]
        if sl["points"] and sl["points"][0][0] == 1:
            sl["points"][0] = (0, sl["points"][0][1])
    return deck


def chunk_slides(slides, max_top=6):
    """Split slides that carry too many points into continuation slides."""
    out = []
    for sl in slides:
        groups = []
        for lvl, t in sl["points"]:
            if lvl == 0 or not groups:
                groups.append([(0, t)])
            else:
                groups[-1].append((1, t))
        n = len(groups)
        if n <= max_top:
            out.append(dict(title=sl["title"], points=sl["points"]))
            continue
        parts = math.ceil(n / max_top)
        size = math.ceil(n / parts)
        for k in range(parts):
            pts = [p for g in groups[k * size:(k + 1) * size] for p in g]
            out.append(dict(title=sl["title"] + ("" if k == 0 else " (cont.)"), points=pts))
    return out


# =============================================================================
#  LOW-LEVEL pptx HELPERS
# =============================================================================
def est_lines(text, size_pt, width_in, cw):
    per_line = max(1, int(width_in * 72 / (size_pt * cw)))
    lines, cur = 1, 0
    for word in text.split():
        L = len(word)
        if cur == 0:
            cur = L
        elif cur + 1 + L <= per_line:
            cur += 1 + L
        else:
            lines += 1
            cur = L
        while cur > per_line:
            lines += 1
            cur -= per_line
    return lines


def split_term(text):
    m = re.match(r"^(.{2,38}?)\s*(?::|\s[-–—]\s)\s+(.+)$", text)
    if m and len(m.group(1).split()) <= 5:
        return m.group(1).strip(), m.group(2).strip()
    return None, text


def set_alpha(shape, pct):
    sf = shape._element.spPr.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf[0]
    for a in clr.findall(qn("a:alpha")):
        clr.remove(a)
    etree.SubElement(clr, qn("a:alpha"), val=str(int(pct * 1000)))


def add_shadow(shape, blur=16, dist=4, alpha=20):
    shape.shadow.inherit = False
    eff = shape._element.spPr.find(qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 12700)),
                          dist=str(int(dist * 12700)), dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def set_gradient(fill, c1, c2, angle):
    fill.gradient()
    fill.gradient_angle = angle
    stops = fill.gradient_stops
    stops[0].color.rgb = rgb(c1)
    stops[0].position = 0
    stops[1].color.rgb = rgb(c2)
    stops[1].position = 1


# =============================================================================
#  PRESENTATION BUILDER
# =============================================================================
SHAPE_FOR = {
    "circles": MSO_SHAPE.OVAL, "rings": MSO_SHAPE.DONUT, "hex": MSO_SHAPE.HEXAGON,
    "triangles": MSO_SHAPE.ISOSCELES_TRIANGLE, "diamonds": MSO_SHAPE.DIAMOND,
    "dots": MSO_SHAPE.OVAL,
}


class Builder:
    def __init__(self, look, ratio="16:9", closing="Thank You", author="",
                 agenda=True, closing_slide=True, progress=None):
        self.look, self.pal, self.st = look, look["pal"], look["style"]
        self.hf, self.bf = look["fonts"]
        self.prs = Presentation()
        self.W, self.H = (10.0, 7.5) if ratio == "4:3" else (13.333, 7.5)
        self.prs.slide_width, self.prs.slide_height = Inches(self.W), Inches(self.H)
        self.mx = 0.75 if self.W > 11 else 0.6
        self.cw_h = FONT_CW.get(self.hf, 0.56) + 0.05
        self.cw_b = FONT_CW.get(self.bf, 0.54) + 0.05
        self.rng = random.Random(look["seed"] + 7)
        self.cycle = [self.pal["primary"], self.pal["secondary"], self.pal["accent"]]
        self.blank = self.prs.slide_layouts[6]
        self.closing_text, self.author = closing, author
        self.use_agenda, self.use_closing = agenda, closing_slide
        self.total, self.deck_title = 1, ""
        self.last_layout = None
        self.progress = progress or (lambda *a: None)

    # ------------------------------------------------------------ primitives
    def new_slide(self):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(self.pal["bg"])
        return s

    def hero_bg(self, s):
        set_gradient(s.background.fill, self.pal["hero1"], self.pal["hero2"], self.st["angle"])

    def shape(self, s, kind, x, y, w, h, fill=None, alpha=None, line=None, line_w=1.0,
              rot=0, radius=None, shadow=False, grad=None, name=None):
        sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if grad:
            set_gradient(sh.fill, grad[0], grad[1], grad[2] if len(grad) > 2 else 0)
        elif fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
            if alpha is not None:
                set_alpha(sh, alpha)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(line_w)
        else:
            sh.line.fill.background()
        if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
        if rot:
            sh.rotation = rot
        if shadow:
            add_shadow(sh)
        else:
            sh.shadow.inherit = False
        if name:
            sh.name = name
        return sh

    def rrect(self, s, x, y, w, h, fill, **kw):
        return self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill,
                          radius=kw.pop("radius", self.st["radius"]), **kw)

    def circle(self, s, cx, cy, d, fill, **kw):
        return self.shape(s, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=fill, **kw)

    def shape_text(self, sh, text, size, color, bold=True, font=None):
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold = Pt(size), bold
        r.font.name = font or self.hf
        r.font.color.rgb = rgb(color)

    def textbox(self, s, x, y, w, h, paras, anchor="t", name=None):
        """paras: list of dict(runs=[(text, dict)], align, after, marL)"""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        for i, pd in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[pd.get("align", "l")]
            p.space_after = Pt(pd.get("after", 0))
            p.line_spacing = pd.get("ls", 1.0)
            if pd.get("marL"):
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(int(Inches(pd["marL"]))))
                pPr.set("indent", "0")
            for text, f in pd["runs"]:
                r = p.add_run()
                r.text = text
                r.font.size = Pt(f["size"])
                r.font.bold = f.get("bold", False)
                r.font.italic = f.get("italic", False)
                r.font.name = f.get("font", self.bf)
                r.font.color.rgb = rgb(f["color"])
        if name:
            tb.name = name
        return tb

    def text(self, s, x, y, w, h, txt, size, color, bold=False, font=None, align="l",
             anchor="t", italic=False, name=None):
        return self.textbox(s, x, y, w, h, [dict(align=align, runs=[(txt, dict(
            size=size, color=color, bold=bold, italic=italic, font=font or self.bf))])],
            anchor=anchor, name=name)

    def fit(self, items, w, h, max_pt, min_pt):
        """items: (text, size_delta, space_after_in, cw) -> largest size that fits."""
        for sz in range(max_pt, min_pt - 1, -1):
            tot = 0.0
            for t, d, after, cw in items:
                tot += est_lines(t, sz + d, w, cw) * (sz + d) * 1.2 / 72 + after
            if tot <= h:
                return sz
        return min_pt

    def fit_title(self, txt, w, h, max_pt=36, min_pt=20):
        return self.fit([(txt, 0, 0, self.cw_h)], w, h, max_pt, min_pt)

    # ------------------------------------------------------------ decoration
    def blob(self, s, cx, cy, d, color, alpha, rot=0):
        kind = self.st["decor"]
        sh = self.shape(s, SHAPE_FOR[kind], cx - d / 2, cy - d / 2, d, d, fill=color, alpha=alpha, rot=rot)
        if kind == "rings":
            sh.adjustments[0] = 0.10
        return sh

    def dot_grid(self, s, x, y, cols, rows, gap, d, color, alpha):
        for r in range(rows):
            for c in range(cols):
                self.circle(s, x + c * gap, y + r * gap, d, color, alpha=alpha)

    def corner_decor(self, s):
        p, W, H = self.pal, self.W, self.H
        a = 16 if p["dark"] else 10
        if self.st["decor"] == "dots":
            self.dot_grid(s, W - 2.1, 0.45, 6, 4, 0.30, 0.09, p["primary"], 45)
            self.dot_grid(s, 0.3, H - 1.5, 4, 4, 0.30, 0.09, p["secondary"], 40)
            self.circle(s, W - 0.3, H - 0.3, 2.4, p["accent"], alpha=a)
        else:
            self.blob(s, W - 0.4, 0.2, 3.6, p["primary"], a, rot=15)
            self.blob(s, W - 2.5, -0.1, 1.5, p["accent"], a + 4, rot=-10)
            self.blob(s, 0.2, H - 0.1, 2.8, p["secondary"], a, rot=10)

    def emblem(self, s, cx, cy, size, main, alt):
        k = self.st["decor"]
        W_ = "FFFFFF"
        if k == "circles":
            for f, col, a in ((1.0, W_, 10), (0.72, alt, 30), (0.46, main, 88), (0.18, W_, 90)):
                self.circle(s, cx, cy, size * f, col, alpha=a)
        elif k == "rings":
            for f, col, a in ((1.0, W_, 25), (0.72, alt, 60), (0.44, main, 95)):
                self.blob(s, cx, cy, size * f, col, a)
        elif k == "dots":
            n = 7
            gap = size / n
            self.dot_grid(s, cx - size / 2 + gap / 2, cy - size / 2 + gap / 2, n, n, gap, gap * 0.28, W_, 35)
            self.circle(s, cx, cy, size * 0.42, main, alpha=92)
            self.circle(s, cx, cy, size * 0.16, W_, alpha=90)
        else:
            for f, col, a, rot in ((1.0, W_, 12, 0), (0.75, alt, 34, 18), (0.48, main, 92, -14)):
                self.blob(s, cx, cy, size * f, col, a, rot=rot)

    def marker(self, s, kind, x, y, d, color, num):
        fg = on_color(color)
        if kind == "num":
            c = self.circle(s, x + d / 2, y + d / 2, d, color)
            self.shape_text(c, str(num), max(11, int(d * 30)), fg, font=self.hf)
        elif kind == "dot":
            self.circle(s, x + d / 2, y + d / 2, d * 0.95, color, alpha=25)
            self.circle(s, x + d / 2, y + d / 2, d * 0.5, color)
        elif kind == "diamond":
            self.shape(s, MSO_SHAPE.DIAMOND, x + d * 0.1, y + d * 0.1, d * 0.8, d * 0.8, fill=color)
        else:  # gem
            self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + d * 0.15, y + d * 0.15, d * 0.7, d * 0.7,
                       fill=color, rot=45, radius=d * 0.18)

    # ------------------------------------------------------------ header/footer
    def header(self, s, title, idx, light_on=None):
        p, W, mx, kind = self.pal, self.W, self.mx, self.st["header"]
        if kind == "bar":
            self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, mx, 0.5, 0.12, 0.95, radius=0.06,
                       grad=(p["primary"], p["accent"], 90))
            x, tw, col = mx + 0.38, W - 2 * mx - 0.38, p["text"]
            y, h = 0.45, 1.05
        elif kind == "underline":
            x, tw, col, y, h = mx, W - 2 * mx, p["text"], 0.38, 1.0
            self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, mx, 1.45, 1.4, 0.09, fill=p["accent"], radius=0.04)
            self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, mx + 1.5, 1.45, 0.3, 0.09, fill=p["primary"], radius=0.04)
        elif kind == "band":
            self.shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, 1.45, grad=(p["hero1"], p["hero2"], self.st["angle"]))
            self.blob(s, W - 1.2, 0.7, 2.6, "FFFFFF", 9, rot=20)
            self.blob(s, W - 2.6, 1.2, 1.2, p["accent"], 55, rot=-15)
            x, tw, col, y, h = mx, W - 2 * mx - 1.6, "FFFFFF", 0.2, 1.05
        else:  # tab
            pill = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, mx, 0.42, 1.25, 0.34, fill=p["primary"], radius=0.17)
            self.shape_text(pill, "TOPIC %02d" % idx, 10, on_color(p["primary"]), font=self.hf)
            x, tw, col, y, h = mx, W - 2 * mx, p["text"], 0.84, 0.85
        sz = self.fit_title(title, tw, h, 34, 20)
        self.text(s, x, y, tw, h, title, sz, col, bold=True, font=self.hf, anchor="m", name="Title")
        return 1.95

    def footer(self, s, idx):
        p, W, H = self.pal, self.W, self.H
        self.shape(s, MSO_SHAPE.RECTANGLE, 0, H - 0.07, W, 0.07, fill=p["muted"], alpha=18)
        self.shape(s, MSO_SHAPE.RECTANGLE, 0, H - 0.07, W * idx / max(1, self.total), 0.07,
                   grad=(p["primary"], p["accent"], 0))
        self.text(s, self.mx, H - 0.5, W * 0.6, 0.28, self.deck_title, 10, p["muted"], name="Footer")
        self.text(s, W - self.mx - 1.2, H - 0.5, 1.2, 0.28, "%d / %d" % (idx, self.total), 10,
                  p["muted"], align="r", name="Page number")

    # ------------------------------------------------------------ special slides
    def slide_title(self, deck):
        s, p, W, H = self.new_slide(), self.pal, self.W, self.H
        self.hero_bg(s)
        v = self.st["title"]
        white, soft = "FFFFFF", mix("FFFFFF", p["hero2"], 0.28)
        label = deck.get("date") or datetime.date.today().strftime("%d %B %Y")
        who = deck.get("author") or self.author
        if v == "split":
            px = W * 0.60
            self.shape(s, MSO_SHAPE.RECTANGLE, px, 0, W - px, H, grad=(p["primary"], p["secondary"], 120))
            self.emblem(s, px + (W - px) / 2, H / 2, min(W - px - 0.5, 4.8), p["accent"], white)
            x, tw = self.mx + 0.1, px - self.mx - 0.7
        elif v == "center":
            self.emblem(s, W - 1.1, 1.0, 3.8, p["accent"], p["secondary"])
            self.emblem(s, 1.0, H - 0.8, 3.0, p["secondary"], p["accent"])
            x, tw = 1.3, W - 2.6
        else:
            if W > 11:
                self.emblem(s, W - 2.8, H / 2, 5.4, p["accent"], p["secondary"])
                x, tw = self.mx + 0.15, W * 0.5
            else:
                self.emblem(s, W - 1.9, H / 2, 3.6, p["accent"], p["secondary"])
                x, tw = self.mx + 0.15, W * 0.46
        al = "c" if v == "center" else "l"
        title = deck["title"] or "Presentation"
        sz = self.fit([(title, 0, 0, self.cw_h)], tw, 2.5, 56, 26)
        lines = est_lines(title, sz, tw, self.cw_h)
        th = lines * sz * 1.2 / 72
        top = 2.35
        # accent pill + label
        px_ = x if al == "l" else (W - 1.2) / 2
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, px_, top - 0.55, 1.2, 0.1, fill=p["accent"], radius=0.05)
        self.text(s, x, top - 0.4, tw, 0.3, label.upper(), 12, soft, bold=True, align=al, font=self.hf)
        self.text(s, x, top, tw, th + 0.1, title, sz, white, bold=True, font=self.hf, align=al, name="Title")
        y = top + th + 0.3
        if deck.get("subtitle"):
            ssz = self.fit([(deck["subtitle"], 0, 0, self.cw_b)], tw, 1.3, 22, 13)
            sh = est_lines(deck["subtitle"], ssz, tw, self.cw_b) * ssz * 1.2 / 72
            self.text(s, x, y, tw, sh + 0.1, deck["subtitle"], ssz, soft, align=al, name="Subtitle")
            y += sh + 0.3
        if who:
            self.shape(s, MSO_SHAPE.RECTANGLE, px_ if al == "l" else (W - 0.6) / 2, H - 1.25, 0.6, 0.04,
                       fill=p["accent"])
            self.text(s, x, H - 1.1, tw, 0.4, who, 15, white, bold=True, align=al, font=self.hf)

    def slide_closing(self):
        s, p, W, H = self.new_slide(), self.pal, self.W, self.H
        self.hero_bg(s)
        self.emblem(s, W - 1.1, 1.0, 3.8, p["accent"], p["secondary"])
        self.emblem(s, 1.1, H - 0.9, 3.2, p["secondary"], p["accent"])
        txt = self.closing_text or "Thank You"
        sz = self.fit([(txt, 0, 0, self.cw_h)], W - 2.4, 1.6, 66, 30)
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, (W - 1.2) / 2, H / 2 - 1.2, 1.2, 0.1, fill=p["accent"], radius=0.05)
        self.text(s, 1.2, H / 2 - 0.95, W - 2.4, 1.5, txt, sz, "FFFFFF", bold=True, font=self.hf,
                  align="c", anchor="m", name="Title")
        self.text(s, 1.2, H / 2 + 0.7, W - 2.4, 0.5, "Questions  &  Discussion", 20,
                  mix("FFFFFF", p["hero2"], 0.28), align="c")
        if self.author:
            self.text(s, 1.2, H - 1.0, W - 2.4, 0.4, self.author, 14, "FFFFFF", bold=True, align="c", font=self.hf)

    def slide_agenda(self, titles, idx):
        s, p, W, H, mx = self.new_slide(), self.pal, self.W, self.H, self.mx
        self.corner_decor(s)
        top = self.header(s, "Agenda", idx)
        n = len(titles)
        cols = 2 if n > 5 else 1
        per = math.ceil(n / cols)
        x0, x1, y1, gap = mx, W - mx, H - 0.95, 0.3
        cw = (x1 - x0 - gap * (cols - 1)) / cols
        if cols == 1:
            cw = min(cw, 8.2)
        ch = min(0.82, (y1 - top - 0.18 * (per - 1)) / per)
        block_h = per * ch + (per - 1) * 0.18
        y0 = top + max(0, (y1 - top - block_h) / 2)
        for i, t in enumerate(titles):
            c, r = divmod(i, per)
            x, y = x0 + c * (cw + gap), y0 + r * (ch + 0.18)
            self.rrect(s, x, y, cw, ch, p["surface"], shadow=True, line=mix(p["bg"], p["muted"], 0.18), line_w=0.75)
            col = self.cycle[i % 3]
            d = ch * 0.64
            nb = self.circle(s, x + 0.22 + d / 2, y + ch / 2, d, col)
            self.shape_text(nb, "%02d" % (i + 1), max(10, int(d * 22)), on_color(col), font=self.hf)
            sz = self.fit([(t, 0, 0, self.cw_b)], cw - d - 0.75, ch - 0.1, 20, 12)
            self.text(s, x + d + 0.5, y, cw - d - 0.7, ch, t, sz, p["text"], bold=True, anchor="m", font=self.hf)
        self.footer(s, idx)

    def slide_section(self, title, idx, sec_no):
        s, p, W, H = self.new_slide(), self.pal, self.W, self.H
        self.hero_bg(s)
        self.emblem(s, W - 3.0, H / 2, 5.0 if W > 11 else 4.0, p["accent"], p["secondary"])
        self.text(s, self.mx, H / 2 - 1.9, 4, 1.4, "%02d" % sec_no, 80, mix("FFFFFF", p["hero2"], 0.45),
                  bold=True, font=self.hf)
        tw = W * 0.55
        sz = self.fit_title(title, tw, 1.8, 46, 24)
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, self.mx, H / 2 - 0.35, 1.2, 0.1, fill=p["accent"], radius=0.05)
        self.text(s, self.mx, H / 2 - 0.1, tw, 1.8, title, sz, "FFFFFF", bold=True, font=self.hf, name="Title")

    # ------------------------------------------------------------ content layouts
    def lay_cards(self, s, d, top, idx):
        p, W, H = self.pal, self.W, self.H
        pts = [t for _, t in d["points"]]
        n = len(pts)
        cols = 3 if (n in (3, 5, 6) and self.W > 11) else (3 if n == 3 else 2)
        rows = math.ceil(n / cols)
        x0, x1, y1, gap = self.mx, W - self.mx, H - 0.95, 0.3
        cw = (x1 - x0 - gap * (cols - 1)) / cols
        ch = min(3.7, (y1 - top - gap * (rows - 1)) / rows)
        y_start = top + max(0, (y1 - top - rows * ch - gap * (rows - 1)) / 2)
        for i, t in enumerate(pts):
            r, c = divmod(i, cols)
            in_row = cols if r < rows - 1 else n - (rows - 1) * cols
            off = (cols - in_row) * (cw + gap) / 2 if r == rows - 1 else 0
            x, y = x0 + off + c * (cw + gap), y_start + r * (ch + gap)
            col = self.cycle[i % 3]
            self.rrect(s, x, y, cw, ch, p["surface"], shadow=True, line=mix(p["bg"], p["muted"], 0.18), line_w=0.75)
            self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.28, y, 0.9, 0.08, fill=col, radius=0.04)
            bd = 0.62
            horiz = ch < 2.7
            if horiz:
                b = self.circle(s, x + 0.3 + bd / 2, y + ch / 2, bd, col)
                tx = x + 0.3 + bd + 0.25
                tw, ty, th = x + cw - 0.3 - tx, y + 0.22, ch - 0.44
            else:
                b = self.circle(s, x + 0.28 + bd / 2, y + 0.3 + bd / 2, bd, col)
                tx, ty, tw = x + 0.28, y + 0.3 + bd + 0.2, cw - 0.56
                th = y + ch - 0.25 - ty
            self.shape_text(b, "%02d" % (i + 1), 14, on_color(col), font=self.hf)
            term, desc = split_term(t)
            if term:
                sz = self.fit([(term, 3, 0.08, self.cw_h), (desc, 0, 0, self.cw_b)], tw, th, 20, 11)
                paras = [dict(after=5, runs=[(term, dict(size=sz + 3, bold=True, color=p["text"], font=self.hf))]),
                         dict(runs=[(desc, dict(size=sz, color=p["muted"] if p["dark"] else mix(p["text"], p["bg"], .25)))])]
            else:
                sz = self.fit([(t, 0, 0, self.cw_b)], tw, th, 22, 11)
                paras = [dict(runs=[(t, dict(size=sz, color=p["text"]))])]
            self.textbox(s, tx, ty, tw, th, paras, anchor="m" if horiz else "t")

    def lay_timeline(self, s, d, top, idx):
        p, W, H = self.pal, self.W, self.H
        pts = [t for _, t in d["points"]]
        n = len(pts)
        x0, x1, y1 = self.mx, W - self.mx, H - 0.95
        step = (x1 - x0) / n
        ly = top + (y1 - top) / 2
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x0, ly - 0.03, x1 - x0, 0.06, grad=(p["primary"], p["accent"], 0), radius=0.03)
        cw = min(step * 1.5, 3.7)
        ch = (y1 - top) / 2 - 0.85
        for i, t in enumerate(pts):
            cx, col = x0 + step * (i + 0.5), self.cycle[i % 3]
            up = (i % 2 == 0)
            cy = ly - 0.62 - ch if up else ly + 0.62
            bx = min(max(cx - cw / 2, x0 - 0.1), x1 - cw + 0.1)
            self.shape(s, MSO_SHAPE.RECTANGLE, cx - 0.015, (cy + ch) if up else (ly + 0.4), 0.03,
                       (ly - 0.4) - (cy + ch) if up else (cy - (ly + 0.4)), fill=col, alpha=60)
            self.circle(s, cx, ly, 0.98, col, alpha=22)
            nd = self.circle(s, cx, ly, 0.68, col)
            self.shape_text(nd, str(i + 1), 18, on_color(col), font=self.hf)
            self.rrect(s, bx, cy, cw, ch, p["surface"], shadow=True, line=mix(p["bg"], p["muted"], 0.18), line_w=0.75)
            term, desc = split_term(t)
            tw, th = cw - 0.5, ch - 0.4
            if term:
                sz = self.fit([(term, 2, 0.06, self.cw_h), (desc, 0, 0, self.cw_b)], tw, th, 18, 11)
                paras = [dict(after=4, runs=[(term, dict(size=sz + 2, bold=True, color=col if not p["dark"] else p["text"], font=self.hf))]),
                         dict(runs=[(desc, dict(size=sz, color=p["text"]))])]
            else:
                sz = self.fit([(t, 0, 0, self.cw_b)], tw, th, 19, 11)
                paras = [dict(runs=[(t, dict(size=sz, color=p["text"]))])]
            self.textbox(s, bx + 0.25, cy + 0.2, tw, th, paras, anchor="m")

    def lay_stats(self, s, d, top, idx, stats):
        p, W, H = self.pal, self.W, self.H
        n = len(stats)
        x0, x1, y1, gap = self.mx, W - self.mx, H - 0.95, 0.35
        cw = (x1 - x0 - gap * (n - 1)) / n
        ch = min(3.9, y1 - top)
        y = top + (y1 - top - ch) / 2
        for i, (num, label) in enumerate(stats):
            x, col = x0 + i * (cw + gap), self.cycle[i % 3]
            self.rrect(s, x, y, cw, ch, p["surface"], shadow=True, line=mix(p["bg"], p["muted"], 0.18), line_w=0.75)
            self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.3, y + 0.3, 0.7, 0.1, fill=col, radius=0.05)
            sz = 64
            while sz > 20 and len(num) * sz * self.cw_h * 1.3 / 72 > cw - 0.6:
                sz -= 2
            self.text(s, x + 0.3, y + 0.55, cw - 0.6, 1.3, num, sz, col, bold=True, font=self.hf, anchor="m")
            self.shape(s, MSO_SHAPE.RECTANGLE, x + 0.3, y + 2.0, cw - 0.6, 0.02, fill=p["muted"], alpha=35)
            lsz = self.fit([(label, 0, 0, self.cw_b)], cw - 0.6, ch - 2.4, 20, 12)
            self.text(s, x + 0.3, y + 2.2, cw - 0.6, ch - 2.4, label, lsz, p["text"])

    def lay_spotlight(self, s, d, top, idx):
        p, W, H = self.pal, self.W, self.H
        x0, x1, y1 = self.mx, W - self.mx, H - 0.95
        h = min(y1 - top, 4.3)
        y = top + (y1 - top - h) / 2
        self.rrect(s, x0, y, x1 - x0, h, p["surface"], shadow=True, line=mix(p["bg"], p["muted"], 0.18), line_w=0.75)
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x0, y, 0.18, h, grad=(p["primary"], p["accent"], 90), radius=0.09)
        self.text(s, x0 + 0.6, y + 0.05, 1.5, 1.6, "\u201C", 120, p["accent"], bold=True, font="Georgia")
        t = d["points"][0][1]
        tw = x1 - x0 - 2.0
        sz = self.fit([(t, 0, 0, self.cw_h)], tw, h - 1.6, 40, 18)
        self.text(s, x0 + 1.3, y + 1.0, tw, h - 1.5, t, sz, p["text"], bold=True, font=self.hf, anchor="m")

    def points_block(self, s, pts, x0, x1, y0, y1, max_pt=24):
        p = self.pal
        items = []
        for lvl, t in pts:
            if lvl == 0 or not items:
                items.append([t, []])
            else:
                items[-1][1].append(t)
        mw = 0.62
        tx = x0 + mw + 0.12
        tw = x1 - tx

        def heights(sz):
            hs = []
            for main, subs in items:
                h = est_lines(main, sz, tw, self.cw_b) * sz * 1.2 / 72
                for sb in subs:
                    h += est_lines(sb, sz - 3, tw - 0.35, self.cw_b) * (sz - 3) * 1.2 / 72 + 0.05
                hs.append(max(h, 0.46))
            return hs

        sz = max_pt
        while sz > 11 and sum(heights(sz)) + 0.28 * (len(items) - 1) > (y1 - y0):
            sz -= 1
        hs = heights(sz)
        free = (y1 - y0) - sum(hs)
        gap = min(0.55, free / (len(items) - 1)) if len(items) > 1 else 0
        y = y0 + max(0, (free - gap * (len(items) - 1)) / 2)
        sub_col = p["muted"]
        for i, ((main, subs), h) in enumerate(zip(items, hs)):
            col = self.cycle[i % 3]
            first_h = sz * 1.2 / 72
            md = 0.46
            self.marker(s, self.st["marker"], x0, y + first_h / 2 - md / 2, md, col, i + 1)
            term, desc = split_term(main)
            if term:
                runs = [(term, dict(size=sz, bold=True, color=col if not p["dark"] else mix(col, "FFFFFF", .25), font=self.hf)),
                        ("  \u2014  " + desc, dict(size=sz, color=p["text"]))]
            else:
                runs = [(main, dict(size=sz, color=p["text"]))]
            paras = [dict(after=0, runs=runs)]
            for sb in subs:
                paras.append(dict(marL=0.3, after=0, runs=[("\u2022  " + sb, dict(size=sz - 3, color=sub_col))]))
            self.textbox(s, tx, y, tw, h + 0.1, paras)
            y += h + gap

    def lay_list(self, s, d, top, idx, sec):
        p, W, H = self.pal, self.W, self.H
        pw = 2.3 if W < 11 else 3.5
        px = W - self.mx - pw
        y1 = H - 0.95
        self.points_block(s, d["points"], self.mx, px - 0.5, top, y1)
        ph = y1 - top
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, px, top, pw, ph, grad=(p["hero1"], p["hero2"], self.st["angle"]),
                   radius=self.st["radius"] + 0.1, shadow=True)
        self.emblem(s, px + pw / 2, top + ph * 0.38, pw * 0.8, p["accent"], p["secondary"])
        self.text(s, px, top + ph - 1.15, pw, 0.9, "%02d" % sec, 44, "FFFFFF", bold=True, font=self.hf, align="c", anchor="m")

    def lay_split(self, s, title, d, idx, sec):
        p, W, H = self.pal, self.W, self.H
        pw = W * 0.34
        self.shape(s, MSO_SHAPE.RECTANGLE, 0, 0, pw, H, grad=(p["hero1"], p["hero2"], 75))
        self.emblem(s, pw * 0.55, H - 1.5, pw * 0.95, p["accent"], p["secondary"])
        self.text(s, 0.5, 0.65, pw - 0.8, 0.9, "%02d" % sec, 44, mix("FFFFFF", p["hero2"], .45), bold=True, font=self.hf)
        tw = pw - 0.9
        sz = self.fit_title(title, tw, 3.2, 34, 18)
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 1.6, 1.0, 0.09, fill=p["accent"], radius=0.045)
        self.text(s, 0.5, 1.85, tw, 3.2, title, sz, "FFFFFF", bold=True, font=self.hf, name="Title")
        a = 14 if p["dark"] else 9
        self.blob(s, W - 0.3, 0.2, 3.2, p["primary"], a, rot=12)
        self.points_block(s, d["points"], pw + 0.7, W - self.mx - 0.2, 0.9, H - 0.95)

    # ------------------------------------------------------------ content dispatcher
    def stats_of(self, pts):
        top = [t for l, t in pts if l == 0]
        if len(top) != len(pts) or not 2 <= len(top) <= 4:
            return None
        res = []
        for t in top:
            m = NUM_RE.search(t)
            if not m or m.start() > 2:
                return None
            label = t[m.end():].strip(" :-\u2013\u2014,.;")
            res.append((m.group(0).strip(), label if len(label) >= 3 else t))
        return res

    def choose_layout(self, d):
        pts = d["points"]
        n_top = sum(1 for l, _ in pts if l == 0)
        has_sub = any(l == 1 for l, _ in pts)
        mx_len = max([len(t) for _, t in pts], default=0)
        if not pts:
            return "section"
        if n_top == 1 and not has_sub and mx_len <= 150:
            return "spotlight"
        if self.stats_of(pts):
            return "stats"
        cand = []
        if not has_sub and 3 <= n_top <= 4 and mx_len <= 130:
            cand += ["cards", "cards"]
        if not has_sub and 5 <= n_top <= 6 and mx_len <= 85 and self.W > 11:
            cand += ["cards"]
        if not has_sub and 3 <= n_top <= 5 and mx_len <= 95:
            cand += ["timeline"]
        cand += ["list", "split"]
        if has_sub:
            cand = ["list", "split"]
        pool = [c for c in cand if c != self.last_layout] or cand
        return self.rng.choice(pool)

    def slide_content(self, d, idx, sec):
        lay = self.choose_layout(d)
        self.last_layout = lay
        if lay == "section":
            self.slide_section(d["title"], idx, sec)
            return lay
        s = self.new_slide()
        if lay == "split":
            self.lay_split(s, d["title"], d, idx, sec)
            self.footer_right(s, idx)
            return lay
        self.corner_decor(s)
        top = self.header(s, d["title"], sec)
        if lay == "cards":
            self.lay_cards(s, d, top, idx)
        elif lay == "timeline":
            self.lay_timeline(s, d, top, idx)
        elif lay == "stats":
            self.lay_stats(s, d, top, idx, self.stats_of(d["points"]))
        elif lay == "spotlight":
            self.lay_spotlight(s, d, top, idx)
        else:
            self.lay_list(s, d, top, idx, sec)
        self.footer(s, idx)
        return lay

    def footer_right(self, s, idx):
        p, W, H = self.pal, self.W, self.H
        pw = W * 0.34
        self.shape(s, MSO_SHAPE.RECTANGLE, pw, H - 0.07, W - pw, 0.07, fill=p["muted"], alpha=18)
        self.shape(s, MSO_SHAPE.RECTANGLE, pw, H - 0.07, (W - pw) * idx / max(1, self.total), 0.07,
                   grad=(p["primary"], p["accent"], 0))
        self.text(s, W - self.mx - 1.2, H - 0.5, 1.2, 0.28, "%d / %d" % (idx, self.total), 10, p["muted"], align="r")

    # ------------------------------------------------------------ build all
    def build(self, deck, out_path):
        slides = chunk_slides(deck["slides"])
        if not slides and not deck["title"]:
            raise ValueError("No content found. Please enter some text first.")
        titles = []
        for sl in slides:
            t = re.sub(r"\s*\(cont\.\)$", "", sl["title"])
            if t not in titles:
                titles.append(t)
        use_agenda = self.use_agenda and len(titles) >= 3
        self.total = 1 + int(use_agenda) + len(slides) + int(self.use_closing)
        self.deck_title = deck["title"] or "Presentation"
        steps = self.total
        done = 0

        def tick():
            nonlocal done
            done += 1
            self.progress(done / steps)

        self.slide_title(deck); tick()
        idx = 1
        if use_agenda:
            idx += 1
            self.slide_agenda(titles[:10], idx); tick()
        sec = 0
        prev_title = None
        for sl in slides:
            idx += 1
            base = re.sub(r"\s*\(cont\.\)$", "", sl["title"])
            if base != prev_title:
                sec += 1
            prev_title = base
            self.slide_content(sl, idx, sec); tick()
        if self.use_closing:
            self.slide_closing(); tick()
        cp = self.prs.core_properties
        cp.title = self.deck_title
        cp.author = deck.get("author") or self.author or "SlideForge"
        self.prs.save(out_path)
        return self.total


def build_presentation(raw_text, out_path, look, ratio="16:9", closing="Thank You", author="",
                       agenda=True, closing_slide=True, progress=None):
    deck = parse_points(raw_text)
    if not deck["title"]:
        deck["title"] = deck["slides"][0]["title"] if deck["slides"] else "Presentation"
    b = Builder(look, ratio, closing, author, agenda, closing_slide, progress)
    return b.build(deck, out_path)


# =============================================================================
#  TKINTER  GUI
# =============================================================================
SAMPLE = """Title: Artificial Intelligence in Education
Subtitle: How smart tools are reshaping the way we learn
Author: Your Name

# What is AI?
- Machine Learning: systems that learn patterns from data
- Natural Language Processing: understanding human language
- Computer Vision: interpreting images and video
- Robotics: machines that act in the physical world

# Why It Matters
- 40% faster grading with automation
- 3x more personalised practice for every student
- 24/7 learning support through smart assistants

# Challenges
- Data privacy
  - student information must be protected
  - clear consent rules are required
- Algorithmic bias
  - models can inherit unfair patterns
- Digital divide
- Teacher training

# Roadmap
- Research: study the needs of students and teachers
- Pilot: test tools in a few classrooms
- Evaluate: measure results and collect feedback
- Scale: roll out across the whole institution

# Key Takeaway
- AI should empower teachers, not replace them
"""

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except ImportError:  # headless environments
    tk = None

UI = dict(bg="#0B1020", panel="#141B31", card="#1B2440", edge="#2A3558", text="#E6EAF7",
          muted="#8A94B8", accent="#6C63FF", accent2="#22D3EE", good="#34D399")


def _hex_lerp(a, b, t):
    ca = [int(a[i:i + 2], 16) for i in (1, 3, 5)]
    cb = [int(b[i:i + 2], 16) for i in (1, 3, 5)]
    return "#%02x%02x%02x" % tuple(round(x + (y - x) * t) for x, y in zip(ca, cb))


if tk:
    class App(tk.Tk):
        def __init__(self):
            super().__init__()
            self.title("SlideForge  -  PPT Generator")
            sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
            self.geometry("%dx%d+20+20" % (min(1220, sw - 60), min(800, sh - 120)))
            self.minsize(900, 520)
            self.configure(bg=UI["bg"])
            self.seed = random.randrange(10 ** 9)
            self.busy = False
            self._styles()
            self._build()
            self.refresh_look()
            self.update_counter()

        # ---------------------------------------------------------------- styles
        def _styles(self):
            st = ttk.Style(self)
            st.theme_use("clam")
            st.configure("TFrame", background=UI["bg"])
            st.configure("Card.TFrame", background=UI["panel"])
            st.configure("TLabel", background=UI["panel"], foreground=UI["text"], font=("Segoe UI", 10))
            st.configure("Muted.TLabel", background=UI["panel"], foreground=UI["muted"], font=("Segoe UI", 9))
            st.configure("H.TLabel", background=UI["panel"], foreground=UI["text"], font=("Segoe UI Semibold", 11))
            st.configure("TCheckbutton", background=UI["panel"], foreground=UI["text"], font=("Segoe UI", 10))
            st.map("TCheckbutton", background=[("active", UI["panel"])])
            st.configure("TRadiobutton", background=UI["panel"], foreground=UI["text"], font=("Segoe UI", 10))
            st.map("TRadiobutton", background=[("active", UI["panel"])])
            st.configure("TCombobox", fieldbackground=UI["card"], background=UI["card"], foreground=UI["text"],
                         arrowcolor=UI["text"], bordercolor=UI["edge"], lightcolor=UI["card"], darkcolor=UI["card"])
            st.map("TCombobox", fieldbackground=[("readonly", UI["card"])], foreground=[("readonly", UI["text"])])
            st.configure("TEntry", fieldbackground=UI["card"], foreground=UI["text"], bordercolor=UI["edge"],
                         insertcolor=UI["text"])
            st.configure("Tool.TButton", background=UI["card"], foreground=UI["text"], borderwidth=0,
                         padding=(12, 6), font=("Segoe UI", 10))
            st.map("Tool.TButton", background=[("active", UI["edge"])])
            st.configure("Roll.TButton", background=UI["accent2"], foreground="#06222B", borderwidth=0,
                         padding=(12, 8), font=("Segoe UI Semibold", 10))
            st.map("Roll.TButton", background=[("active", "#67E8F9")])
            st.configure("Go.TButton", background=UI["accent"], foreground="white", borderwidth=0,
                         padding=(18, 12), font=("Segoe UI Semibold", 13))
            st.map("Go.TButton", background=[("active", "#8077FF"), ("disabled", "#3B4170")])
            st.configure("Horizontal.TProgressbar", troughcolor=UI["card"], background=UI["good"],
                         bordercolor=UI["card"], lightcolor=UI["good"], darkcolor=UI["good"])

        # ---------------------------------------------------------------- layout
        def _build(self):
            # banner with gradient
            self.banner = tk.Canvas(self, height=92, highlightthickness=0, bg=UI["bg"])
            self.banner.pack(fill="x")
            self.banner.bind("<Configure>", self._draw_banner)

            bottom = ttk.Frame(self)
            bottom.pack(side="bottom", fill="x", padx=18, pady=(4, 12))

            main = ttk.Frame(self)
            main.pack(fill="both", expand=True, padx=18, pady=(12, 6))
            main.columnconfigure(0, weight=1)
            main.columnconfigure(1, weight=0)
            main.rowconfigure(0, weight=1)

            # ----- left: text editor
            left = ttk.Frame(main, style="Card.TFrame")
            left.grid(row=0, column=0, sticky="nsew", padx=(0, 14))
            left.columnconfigure(0, weight=1)
            left.rowconfigure(2, weight=1)
            top = ttk.Frame(left, style="Card.TFrame")
            top.grid(row=0, column=0, sticky="ew", padx=14, pady=(12, 6))
            ttk.Label(top, text="Your presentation points", style="H.TLabel").pack(side="left")
            for txt, cmd in (("Clear", self.clear), ("Sample", self.load_sample), ("Open .txt", self.open_txt)):
                ttk.Button(top, text=txt, style="Tool.TButton", command=cmd).pack(side="right", padx=(6, 0))
            ttk.Label(left, text="Title:  Subtitle:  Author:   |   # Slide heading   |   - bullet   |   indent = sub-point   |   "
                                 "Term: description  -> bold term",
                      style="Muted.TLabel").grid(row=1, column=0, sticky="w", padx=14)
            box = tk.Frame(left, bg=UI["edge"], bd=0)
            box.grid(row=2, column=0, sticky="nsew", padx=14, pady=8)
            box.rowconfigure(0, weight=1)
            box.columnconfigure(0, weight=1)
            self.txt = tk.Text(box, wrap="word", undo=True, bg="#0E1428", fg=UI["text"], insertbackground="white",
                               selectbackground=UI["accent"], relief="flat", padx=12, pady=10,
                               font=("Consolas", 12), spacing1=2, spacing3=2)
            sb = ttk.Scrollbar(box, command=self.txt.yview)
            self.txt.configure(yscrollcommand=sb.set)
            self.txt.grid(row=0, column=0, sticky="nsew", padx=1, pady=1)
            sb.grid(row=0, column=1, sticky="ns", pady=1)
            self.txt.bind("<KeyRelease>", lambda e: self.update_counter())
            self.txt.tag_configure("ph", foreground=UI["muted"])
            self.counter = ttk.Label(left, text="", style="Muted.TLabel")
            self.counter.grid(row=3, column=0, sticky="w", padx=14, pady=(0, 10))

            # ----- right: options
            outer = ttk.Frame(main, style="Card.TFrame", width=350)
            outer.grid(row=0, column=1, sticky="ns")
            outer.grid_propagate(False)
            outer.rowconfigure(0, weight=1)
            rc = tk.Canvas(outer, width=334, bg=UI["panel"], highlightthickness=0)
            rsb = ttk.Scrollbar(outer, orient="vertical", command=rc.yview)
            rc.configure(yscrollcommand=rsb.set)
            rc.grid(row=0, column=0, sticky="ns")
            rsb.grid(row=0, column=1, sticky="ns")
            right = ttk.Frame(rc, style="Card.TFrame")
            rwin = rc.create_window((0, 0), window=right, anchor="nw", width=334)
            right.bind("<Configure>", lambda e: rc.configure(scrollregion=rc.bbox("all")))
            rc.bind("<Enter>", lambda e: rc.bind_all("<MouseWheel>", lambda ev: rc.yview_scroll(int(-ev.delta / 120), "units")))
            rc.bind("<Leave>", lambda e: rc.unbind_all("<MouseWheel>"))
            pad = dict(padx=16, anchor="w")
            ttk.Label(right, text="Design studio", style="H.TLabel").pack(pady=(14, 8), **pad)

            ttk.Label(right, text="Colour theme").pack(**pad)
            self.theme = tk.StringVar(value=list(PRESETS)[0])
            cb = ttk.Combobox(right, textvariable=self.theme, values=list(PRESETS), state="readonly", width=38)
            cb.pack(padx=16, pady=(2, 8), fill="x")
            cb.bind("<<ComboboxSelected>>", lambda e: self.refresh_look())

            row = ttk.Frame(right, style="Card.TFrame")
            row.pack(padx=16, pady=2, fill="x")
            ttk.Label(row, text="Mode").pack(side="left")
            self.mode = tk.StringVar(value="auto")
            for v, t in (("auto", "Auto"), ("light", "Light"), ("dark", "Dark")):
                ttk.Radiobutton(row, text=t, value=v, variable=self.mode, command=self.refresh_look).pack(side="left", padx=(10, 0))

            row = ttk.Frame(right, style="Card.TFrame")
            row.pack(padx=16, pady=2, fill="x")
            ttk.Label(row, text="Slide size").pack(side="left")
            self.ratio = tk.StringVar(value="16:9")
            for v in ("16:9", "4:3"):
                ttk.Radiobutton(row, text=v, value=v, variable=self.ratio).pack(side="left", padx=(10, 0))

            ttk.Label(right, text="Fonts").pack(pady=(8, 0), **pad)
            self.font = tk.StringVar(value=list(FONT_PAIRS)[0])
            fb = ttk.Combobox(right, textvariable=self.font, values=list(FONT_PAIRS), state="readonly", width=38)
            fb.pack(padx=16, pady=(2, 8), fill="x")
            fb.bind("<<ComboboxSelected>>", lambda e: self.refresh_look())

            self.agenda = tk.BooleanVar(value=True)
            self.closing = tk.BooleanVar(value=True)
            self.autoroll = tk.BooleanVar(value=True)
            ttk.Checkbutton(right, text="Add agenda slide", variable=self.agenda).pack(**pad)
            ttk.Checkbutton(right, text="Add closing slide", variable=self.closing).pack(**pad)
            ttk.Checkbutton(right, text="Brand-new look after every export", variable=self.autoroll).pack(**pad)

            ttk.Label(right, text="Presenter name").pack(pady=(8, 0), **pad)
            self.author = tk.StringVar()
            ttk.Entry(right, textvariable=self.author).pack(padx=16, pady=(2, 4), fill="x")
            ttk.Label(right, text="Closing text").pack(**pad)
            self.closing_txt = tk.StringVar(value="Thank You")
            ttk.Entry(right, textvariable=self.closing_txt).pack(padx=16, pady=(2, 8), fill="x")

            ttk.Label(right, text="Current look (preview)", style="H.TLabel").pack(pady=(6, 4), **pad)
            self.sw = tk.Canvas(right, height=44, width=316, bg=UI["panel"], highlightthickness=0)
            self.sw.pack(padx=16)
            self.look_lbl = ttk.Label(right, text="", style="Muted.TLabel", wraplength=310, justify="left")
            self.look_lbl.pack(padx=16, pady=(4, 6), anchor="w")
            ttk.Button(right, text="\U0001F3B2  Roll a new look", style="Roll.TButton", command=self.roll).pack(padx=16, pady=(0, 12), fill="x")

            # ----- bottom bar
            bottom.columnconfigure(1, weight=1)
            self.go = ttk.Button(bottom, text="\u2728  Generate PowerPoint", style="Go.TButton", command=self.generate)
            self.go.grid(row=0, column=0, rowspan=2, padx=(0, 16))
            self.bar = ttk.Progressbar(bottom, mode="determinate", maximum=100)
            self.bar.grid(row=0, column=1, sticky="ew", pady=(6, 2))
            self.status = tk.Label(bottom, text="Ready - paste your points and press Generate.", bg=UI["bg"],
                                   fg=UI["muted"], font=("Segoe UI", 10), anchor="w")
            self.status.grid(row=1, column=1, sticky="ew")

        def _draw_banner(self, _=None):
            c = self.banner
            c.delete("all")
            w = max(c.winfo_width(), 800)
            for x in range(0, w, 3):
                c.create_rectangle(x, 0, x + 3, 92, fill=_hex_lerp("#4F46E5", "#06B6D4", x / w), outline="")
            for cx, cy, r, col in ((w - 80, 20, 70, "#FFFFFF"), (w - 190, 90, 45, "#FFFFFF"), (60, 95, 60, "#FFFFFF")):
                c.create_oval(cx - r, cy - r, cx + r, cy + r, outline=col, width=2, stipple="gray25")
            c.create_text(28, 34, anchor="w", text="SlideForge", fill="white", font=("Segoe UI Black", 26))
            c.create_text(30, 66, anchor="w", fill="#E0F2FE", font=("Segoe UI", 11),
                          text="Turn plain points into a designer-grade presentation - a fresh look every time.")

        # ---------------------------------------------------------------- actions
        def text_value(self):
            return self.txt.get("1.0", "end").strip()

        def update_counter(self):
            raw = self.text_value()
            words = len(raw.split())
            try:
                deck = parse_points(raw)
                n = len(chunk_slides(deck["slides"])) + 1
            except Exception:
                n = 0
            self.counter.config(text="%d words   |   about %d content slides" % (words, max(0, n - 1)))

        def clear(self):
            self.txt.delete("1.0", "end")
            self.update_counter()

        def load_sample(self):
            self.txt.delete("1.0", "end")
            self.txt.insert("1.0", SAMPLE)
            self.update_counter()

        def open_txt(self):
            path = filedialog.askopenfilename(filetypes=[("Text files", "*.txt *.md"), ("All files", "*.*")])
            if path:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    self.txt.delete("1.0", "end")
                    self.txt.insert("1.0", f.read())
                self.update_counter()

        def current_look(self):
            f = self.font.get()
            return make_look(self.seed, self.theme.get(), self.mode.get(), f)

        def refresh_look(self):
            look = self.current_look()
            p = look["pal"]
            self.sw.delete("all")
            cols = [p["bg"], p["surface"], p["text"], p["primary"], p["secondary"], p["accent"], p["hero1"], p["hero2"]]
            w = 316 / len(cols)
            for i, c in enumerate(cols):
                self.sw.create_rectangle(i * w + 1, 2, (i + 1) * w - 1, 42, fill="#" + c, outline=UI["edge"])
            self.look_lbl.config(text=describe_look(look))

        def roll(self):
            self.seed = random.randrange(10 ** 9)
            self.refresh_look()

        def set_status(self, msg, color=None):
            self.status.config(text=msg, fg=color or UI["muted"])

        def generate(self):
            if self.busy:
                return
            raw = self.text_value()
            if not raw:
                messagebox.showwarning("Nothing to build", "Please type or paste your presentation points first.")
                return
            default = "Presentation_%s.pptx" % datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            path = filedialog.asksaveasfilename(defaultextension=".pptx", initialfile=default,
                                                filetypes=[("PowerPoint", "*.pptx")])
            if not path:
                return
            look = self.current_look()
            args = dict(ratio=self.ratio.get(), closing=self.closing_txt.get().strip() or "Thank You",
                        author=self.author.get().strip(), agenda=self.agenda.get(),
                        closing_slide=self.closing.get())
            self.busy = True
            self.go.state(["disabled"])
            self.bar["value"] = 0
            self.set_status("Designing your slides ...", UI["accent2"])

            def work():
                try:
                    total = build_presentation(raw, path, look, progress=lambda f: self.after(0, self._prog, f), **args)
                    self.after(0, self._done, path, total, None)
                except Exception as e:  # noqa
                    self.after(0, self._done, path, 0, e)

            threading.Thread(target=work, daemon=True).start()

        def _prog(self, f):
            self.bar["value"] = f * 100

        def _done(self, path, total, err):
            self.busy = False
            self.go.state(["!disabled"])
            if err:
                self.set_status("Failed: %s" % err, "#F87171")
                messagebox.showerror("Could not create PPT", str(err))
                return
            self.bar["value"] = 100
            self.set_status("Saved %d slides  ->  %s" % (total, path), UI["good"])
            if self.autoroll.get():
                self.roll()
            if messagebox.askyesno("Done!", "Presentation created with %d slides.\n\nOpen it now?" % total):
                try:
                    if sys.platform.startswith("win"):
                        os.startfile(path)  # type: ignore[attr-defined]
                    elif sys.platform == "darwin":
                        subprocess.Popen(["open", path])
                    else:
                        subprocess.Popen(["xdg-open", path])
                except Exception as e:  # noqa
                    messagebox.showinfo("Open file", "Saved at:\n%s\n\n(%s)" % (path, e))


def main():
    if tk is None:
        print("Tkinter is not installed. On Linux: sudo apt install python3-tk")
        sys.exit(1)
    App().mainloop()


if __name__ == "__main__":
    main()
